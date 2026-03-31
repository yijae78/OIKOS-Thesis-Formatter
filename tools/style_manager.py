"""Style manager — view and apply paragraph styles."""

from docx import Document
from docx.oxml.ns import qn
from pathlib import Path


def get_all_styles(doc: Document) -> list[dict]:
    """Get all defined styles in the document."""
    styles = []
    for style in doc.styles:
        if style.type is not None:
            styles.append({
                'name': style.name,
                'style_id': style.style_id,
                'type': str(style.type),
                'builtin': style.builtin,
                'base_style': style.base_style.name if style.base_style else None,
            })
    return styles


def get_paragraph_styles(doc: Document) -> list[dict]:
    """Get only paragraph styles (not character, table, etc.)."""
    from docx.enum.style import WD_STYLE_TYPE
    return [s for s in get_all_styles(doc) if s['type'] == str(WD_STYLE_TYPE.PARAGRAPH)]


def get_style_usage(doc: Document) -> dict[str, list[int]]:
    """Map style names to paragraph indices where they're used."""
    usage: dict[str, list[int]] = {}
    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "(No Style)"
        if style_name not in usage:
            usage[style_name] = []
        usage[style_name].append(i)
    return usage


def change_paragraph_style(doc: Document, paragraph_index: int, new_style_name: str) -> None:
    """Change the style of a specific paragraph."""
    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        raise ValueError(f"Paragraph index {paragraph_index} out of range (0-{len(doc.paragraphs) - 1})")

    # Verify style exists
    try:
        style = doc.styles[new_style_name]
    except KeyError:
        available = [s.name for s in doc.styles]
        raise ValueError(f"Style '{new_style_name}' not found. Available: {available}")

    para = doc.paragraphs[paragraph_index]
    old_style = para.style.name if para.style else "(No Style)"
    para.style = style
    print(f"[OK] Paragraph {paragraph_index}: '{old_style}' → '{new_style_name}'")
    print(f"     Text: {para.text[:80]}...")


def change_style_for_range(doc: Document, start_idx: int, end_idx: int, new_style_name: str) -> int:
    """Change style for a range of paragraphs. Returns count changed."""
    try:
        doc.styles[new_style_name]
    except KeyError:
        raise ValueError(f"Style '{new_style_name}' not found")

    count = 0
    for i in range(start_idx, min(end_idx + 1, len(doc.paragraphs))):
        doc.paragraphs[i].style = doc.styles[new_style_name]
        count += 1

    print(f"[OK] Changed {count} paragraphs (index {start_idx}-{end_idx}) to style '{new_style_name}'")
    return count


def find_paragraphs_by_style(doc: Document, style_name: str) -> list[dict]:
    """Find all paragraphs with a given style."""
    results = []
    for i, para in enumerate(doc.paragraphs):
        if para.style and para.style.name == style_name:
            results.append({
                'index': i,
                'text': para.text[:100] + ('...' if len(para.text) > 100 else ''),
            })
    return results


def get_style_font_info(doc: Document, style_name: str) -> dict:
    """Get font information for a style."""
    try:
        style = doc.styles[style_name]
    except KeyError:
        return {'error': f"Style '{style_name}' not found"}

    font = style.font
    pf = style.paragraph_format

    info = {
        'name': style_name,
        'font_name': font.name,
        'font_size': str(font.size) if font.size else None,
        'bold': font.bold,
        'italic': font.italic,
        'color': str(font.color.rgb) if font.color and font.color.rgb else None,
    }

    if pf:
        info.update({
            'alignment': str(pf.alignment) if pf.alignment else None,
            'space_before': str(pf.space_before) if pf.space_before else None,
            'space_after': str(pf.space_after) if pf.space_after else None,
            'line_spacing': str(pf.line_spacing) if pf.line_spacing else None,
            'first_line_indent': str(pf.first_line_indent) if pf.first_line_indent else None,
            'left_indent': str(pf.left_indent) if pf.left_indent else None,
        })

    return info


def print_style_report(doc: Document) -> None:
    """Print a formatted report of all styles and their usage."""
    usage = get_style_usage(doc)

    print(f"\n{'─' * 60}")
    print("STYLE USAGE REPORT")
    print(f"{'─' * 60}")
    print(f"{'Count':>6} {'Style Name':<30} {'First Paragraph Text'}")
    print(f"{'─' * 6} {'─' * 30} {'─' * 30}")

    for style_name, indices in sorted(usage.items(), key=lambda x: -len(x[1])):
        first_text = doc.paragraphs[indices[0]].text[:40] if indices else ""
        print(f"{len(indices):>6} {style_name:<30} {first_text}")
