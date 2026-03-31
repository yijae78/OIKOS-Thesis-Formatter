"""Header/footer manager — edit headers, footers, and their linkage per section."""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path


def get_header_footer_map(doc: Document) -> list[dict]:
    """Get detailed header/footer info for each section."""
    result = []
    for i, section in enumerate(doc.sections):
        sec_info = {
            'section': i + 1,
            'different_first_page': section.different_first_page_header_footer,
            'headers': {},
            'footers': {},
            'linked_to_previous': {},
        }

        # Check header linkage
        for hf_type, attr_default, attr_first, attr_even in [
            ('header', 'header', 'first_page_header', 'even_page_header'),
            ('footer', 'footer', 'first_page_footer', 'even_page_footer'),
        ]:
            for variant, attr_name in [('default', attr_default), ('first_page', attr_first), ('even_page', attr_even)]:
                try:
                    obj = getattr(section, attr_name)
                    is_linked = obj.is_linked_to_previous
                    text_parts = [p.text for p in obj.paragraphs if p.text.strip()]
                    text = ' | '.join(text_parts) if text_parts else '(empty)'

                    sec_info[f'{hf_type}s'][variant] = {
                        'text': text,
                        'linked_to_previous': is_linked,
                        'paragraph_count': len(obj.paragraphs),
                    }
                    sec_info['linked_to_previous'][f'{hf_type}_{variant}'] = is_linked
                except Exception:
                    pass

        result.append(sec_info)
    return result


def set_header_text(doc: Document, section_index: int, text: str, variant: str = 'default') -> None:
    """Set header text for a section.

    Args:
        section_index: 0-based section index
        text: Text to set (replaces all existing header content)
        variant: 'default', 'first_page', or 'even_page'
    """
    if section_index < 0 or section_index >= len(doc.sections):
        raise ValueError(f"Section index {section_index} out of range")

    section = doc.sections[section_index]
    attr_map = {
        'default': 'header',
        'first_page': 'first_page_header',
        'even_page': 'even_page_header',
    }

    if variant not in attr_map:
        raise ValueError(f"Invalid variant '{variant}'. Valid: {list(attr_map.keys())}")

    header = getattr(section, attr_map[variant])

    # Unlink from previous if needed
    if header.is_linked_to_previous:
        header.is_linked_to_previous = False

    # Clear existing content and set new text
    for p in header.paragraphs:
        p.clear()

    if header.paragraphs:
        header.paragraphs[0].text = text
    else:
        header.add_paragraph(text)

    print(f"[OK] Section {section_index + 1} {variant} header: '{text[:50]}'")


def set_footer_text(doc: Document, section_index: int, text: str, variant: str = 'default') -> None:
    """Set footer text for a section."""
    if section_index < 0 or section_index >= len(doc.sections):
        raise ValueError(f"Section index {section_index} out of range")

    section = doc.sections[section_index]
    attr_map = {
        'default': 'footer',
        'first_page': 'first_page_footer',
        'even_page': 'even_page_footer',
    }

    if variant not in attr_map:
        raise ValueError(f"Invalid variant '{variant}'. Valid: {list(attr_map.keys())}")

    footer = getattr(section, attr_map[variant])

    if footer.is_linked_to_previous:
        footer.is_linked_to_previous = False

    for p in footer.paragraphs:
        p.clear()

    if footer.paragraphs:
        footer.paragraphs[0].text = text
    else:
        footer.add_paragraph(text)

    print(f"[OK] Section {section_index + 1} {variant} footer: '{text[:50]}'")


def add_page_number_to_footer(doc: Document, section_index: int, variant: str = 'default') -> None:
    """Add an automatic page number field to a section's footer.

    This inserts a PAGE field code that Word will render as the current page number.
    """
    if section_index < 0 or section_index >= len(doc.sections):
        raise ValueError(f"Section index {section_index} out of range")

    section = doc.sections[section_index]
    attr_map = {
        'default': 'footer',
        'first_page': 'first_page_footer',
        'even_page': 'even_page_footer',
    }

    footer = getattr(section, attr_map.get(variant, 'footer'))

    if footer.is_linked_to_previous:
        footer.is_linked_to_previous = False

    # Clear existing
    for p in footer.paragraphs:
        p.clear()

    # Add PAGE field
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()

    # Center alignment
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add PAGE field code
    run = para.add_run()
    fld_char_begin = OxmlElement('w:fldChar')
    fld_char_begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(fld_char_begin)

    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    run._r.append(instr_text)

    fld_char_end = OxmlElement('w:fldChar')
    fld_char_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char_end)

    print(f"[OK] Section {section_index + 1} {variant} footer: Added page number field")


def link_to_previous(doc: Document, section_index: int, hf_type: str = 'both') -> None:
    """Link a section's headers/footers to the previous section.

    Args:
        section_index: 0-based section index
        hf_type: 'header', 'footer', or 'both'
    """
    if section_index < 1 or section_index >= len(doc.sections):
        raise ValueError(f"Section {section_index} cannot link to previous (must be >= 1)")

    section = doc.sections[section_index]

    if hf_type in ('header', 'both'):
        section.header.is_linked_to_previous = True
        print(f"[OK] Section {section_index + 1}: Header linked to previous")

    if hf_type in ('footer', 'both'):
        section.footer.is_linked_to_previous = True
        print(f"[OK] Section {section_index + 1}: Footer linked to previous")


def unlink_from_previous(doc: Document, section_index: int, hf_type: str = 'both') -> None:
    """Unlink a section's headers/footers from the previous section."""
    if section_index < 1 or section_index >= len(doc.sections):
        raise ValueError(f"Section {section_index} out of range")

    section = doc.sections[section_index]

    if hf_type in ('header', 'both'):
        section.header.is_linked_to_previous = False
        print(f"[OK] Section {section_index + 1}: Header unlinked from previous")

    if hf_type in ('footer', 'both'):
        section.footer.is_linked_to_previous = False
        print(f"[OK] Section {section_index + 1}: Footer unlinked from previous")


def set_different_first_page(doc: Document, section_index: int, enabled: bool) -> None:
    """Enable or disable 'Different First Page' for a section."""
    if section_index < 0 or section_index >= len(doc.sections):
        raise ValueError(f"Section index {section_index} out of range")

    section = doc.sections[section_index]
    section.different_first_page_header_footer = enabled
    state = "enabled" if enabled else "disabled"
    print(f"[OK] Section {section_index + 1}: Different first page {state}")


def print_header_footer_report(doc: Document) -> None:
    """Print formatted header/footer report."""
    hf_map = get_header_footer_map(doc)

    print(f"\n{'─' * 70}")
    print("HEADER / FOOTER REPORT")
    print(f"{'─' * 70}")

    for sec in hf_map:
        print(f"\n  Section {sec['section']} (Different first page: {sec['different_first_page']})")

        for hf_type in ['headers', 'footers']:
            for variant, info in sec[hf_type].items():
                linked = "→LINKED" if info['linked_to_previous'] else "CUSTOM"
                text = info['text'][:40]
                print(f"    {hf_type[:-1]:>6} ({variant:<11}): [{linked:>7}] {text}")
