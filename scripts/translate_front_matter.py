"""Translate all Korean front matter (P0~P227) in thesis_english.docx.
Uses XML-level paragraph clearing to handle TOC hyperlinks and field codes."""
import re
import os
import shutil
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn

def clear_paragraph_content(para, new_text, font_name='Times New Roman'):
    """Clear ALL content (runs, hyperlinks, fields) and set new text."""
    elem = para._element
    pPr = elem.find(qn('w:pPr'))

    # Remove all children except pPr
    for child in list(elem):
        if child.tag != qn('w:pPr'):
            elem.remove(child)

    # Add new run with text
    run = para.add_run(new_text)
    run.font.name = font_name

def translate_front_matter():
    base = os.path.join(os.path.dirname(__file__), '..')
    working = os.path.join(base, 'output', 'thesis_english.docx')

    # Backup
    backup_dir = os.path.join(base, 'backups')
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup = os.path.join(backup_dir, f'thesis_pre_fm_{ts}.docx')
    shutil.copy2(working, backup)
    print(f"Backup: {os.path.basename(backup)}")

    doc = Document(working)

    # All front matter translations indexed by paragraph position
    translations = {
        # === Korean Abstract (P17-P26) ===
        17: "ABSTRACT",
        18: "Yijae Shin",
        19: '2026\t\t\t"A Study on the Mechanism by Which AI Utilization Training Requests Fail to Translate into Organizational Performance \u2013 A Seizing-Stage Analysis Based on Dynamic Capabilities Theory"',
        20: "Despite repeated observations that judgments formed within organizations after AI utilization training requests fail to convert into strategic decision-making and tangible performance outcomes, existing research has predominantly reduced this phenomenon to issues of individual-level technology acceptance or training effectiveness. However, such approaches fail to structurally explain under what processual and institutional conditions approved judgments within organizations fail to convert into formal decision-making.",
        21: "A qualitative research design centered on a single case was adopted to conduct an in-depth analysis of the pathway through which approved judgments failed to be placed on strategic agendas. Through process tracing and pattern matching, the study empirically reconstructed whether decision-making was initiated after judgment formation, the flow of procedural progress, and the conditions of institutional stagnation.",
        22: "Case analysis revealed that even in the absence of explicit opposition, when failure of agenda-setting, ambiguity of responsible parties, and the absence of institutional decision-making triggers combine, judgments become entrenched in a state of non-decision. Such conversion failure was analyzed as possessing inherently self-reinforcing properties at the processual level.",
        23: "The observed empirical evidence clearly reveals the limitations of conventional perspectives. In contrast to prior approaches that regarded Seizing as a matter of simple opportunity execution or automatic strategic choice, Seizing can be reconceptualized as processual infrastructure through which organizational judgments convert into decisions.",
        24: "This study explores the structural conditions under which judgments conceived within organizations float without transitioning into decision-making, thereby attempting a processual extension of the Seizing stage within Dynamic Capabilities Theory and contributing to the theoretical formalization of the generation, legitimation, agenda-setting, and dissipation processes of organizational judgment through the concept of the Judgment Life Cycle.",
        26: "Keywords: AI Utilization Training Requests, Dynamic Capabilities Theory, Seizing, Judgment Life Cycle, Conditional Trigger Model, Non-decision, Non-event",

        # === Section headings ===
        27: "TABLE OF CONTENTS",
        159: "LIST OF TABLES",
        182: "LIST OF FIGURES",
        206: "LIST OF ABBREVIATIONS",
        212: "DEDICATION",
        218: "ACKNOWLEDGEMENTS",

        # === TOC title entries (P28-P29) ===
        29: '\u300cA Study on the Mechanism by Which AI Utilization Training Requests Fail to Translate into Organizational Performance \u2013 A Seizing-Stage Analysis Based on Dynamic Capabilities Theory\u300d',

        # === TOC front matter entries ===
        31: "ABSTRACT\tiv",
        32: "TABLE OF CONTENTS\tvi",
        33: "LIST OF TABLES\txi",
        34: "LIST OF FIGURES\txii",
        35: "LIST OF ABBREVIATIONS\txiii",
        36: "DEDICATION\txiv",
        37: "ACKNOWLEDGEMENTS\txv",

        # === TOC Chapter 1 ===
        38: "CHAPTER 1  INTRODUCTION\t1",
        39: "1.1. Background of the Study\t1",
        40: "1.2. Thesis Statement\t3",
        41: "1.3. Purpose of the Study\t3",
        42: "1.4. Goals of the Study\t5",
        43: "1.5. Significance of the Study\t6",
        44: "1.5.1. Theoretical Significance\t7",
        45: "1.5.2. Practical Significance\t8",
        46: "1.5.3. Policy and Societal Significance\t8",
        47: "1.6. Central Research Issue\t9",
        48: "1.7. Research Questions\t11",
        49: "1.8. Delimitations and Limitations of the Study\t14",
        50: "1.9. Definitions\t17",
        51: "1.9.1. AI Utilization Training\t18",
        52: "1.9.2. Organizational Performance\t19",
        53: "1.9.3. Judgment\t19",
        54: "1.9.4. Decision-Making\t20",
        55: "1.9.5. Micro-level Mechanism\t20",
        56: "1.9.6. Dynamic Capabilities and Seizing\t21",
        57: "1.9.7. Theoretical Generalization\t21",
        # P58 already English
        59: "1.9.9. Seizing Stagnation (or Seizing Non-initiation)\t23",
        # P60 already English
        61: "1.10. Analytical Propositions\t24",
        62: "1.10.1. Analytical Proposition Regarding Post-Request Judgment Formation (P1)\t24",
        63: "1.10.2. Analytical Proposition Regarding Judgment\u2013Decision-Making Conversion Failure (P2)\t25",
        64: "1.10.3. Analytical Proposition Regarding Micro-level Mechanisms at the Seizing Stage (P3)\t25",
        65: "1.11. Methodology\t26",
        66: "1.12. Procedures\t26",
        67: "1.13. Summary\t27",

        # === TOC Chapter 2 ===
        68: "CHAPTER 2  LITERATURE REVIEW\t30",
        69: "2.1. Theoretical Development of Dynamic Capabilities Theory\t30",
        70: "2.2. The Concept of Seizing and Its Microfoundations\t32",
        71: "2.3. Case-based Theory Building and Single-Case Approach\t34",
        72: "2.4. Prior Research on AI Adoption and Organizational Transformation\t37",
        73: "2.4.1. Research on AI Adoption as a Catalyst for Organizational Transformation\t37",
        74: "2.4.2. Research on AI Utilization Capabilities and Organizational Performance\t38",
        75: "2.5. The Knowing\u2013Doing Gap and Research Gap in Seizing Mechanisms\t40",
        76: "2.6. Summary\t43",

        # === TOC Chapter 3 ===
        77: "CHAPTER 3  RESEARCH METHODOLOGY\t46",
        78: "3.1. Research Design Overview\t46",
        79: "3.2. Qualitative Research Approach and Single-Case Study Design\t48",
        80: "3.3. Case Selection Rationale\t51",
        81: "3.4. Data Collection\t54",
        82: "3.4.1. Coding and Interpretation Criteria\t55",
        83: "3.4.2. Decision Conversion Checklist\t56",
        84: "3.4.3. Basic Principles of Data Collection\t57",
        85: "3.4.4. Official Document Data\t57",
        86: "3.4.5. Electronic Approval System Records\t58",
        87: "3.4.6. Non-response Status and Procedural Gap Data\t58",
        88: "3.4.7. Participant Observation Data\t59",
        89: "3.4.8. Timing and Scope of Data Collection\t59",
        90: "3.5. Data Analysis\t60",
        91: "3.5.1. Overview of Analytical Strategy\t61",
        92: "3.5.2. Application of Process Tracing\t63",
        93: "3.5.3. Theory\u2013Case Comparison through Pattern Matching\t64",
        94: "3.5.4. Linking Explanatory Hypotheses and Empirical Data\t66",
        95: "3.5.5. Ensuring Analytical Reliability\t67",
        96: "3.6. Trustworthiness and Rigor in Qualitative Research\t67",
        97: "3.6.1. Concepts of Reliability and Validity in Qualitative Research\t67",
        98: "3.6.2. Credibility Assurance Strategies\t68",
        99: "3.6.3. Considerations for Transferability\t69",
        100: "3.6.4. Dependability Assurance Strategies\t70",
        101: "3.6.5. Confirmability Assurance Strategies\t70",
        102: "3.6.6. Validity of Non-event Analysis\t71",
        103: "3.7. Research Ethics and Positionality\t71",
        104: "3.7.1. Research Ethics Principles and Data Protection\t72",
        105: "3.7.2. Considerations Regarding IRB Review\t73",
        106: "3.7.3. Insider Positionality\t73",
        107: "3.7.4. Insider Bias Control Procedures\t74",
        108: "3.7.5. Summary\t75",
        109: "3.8. Summary\t75",

        # === TOC Chapter 4 ===
        110: "CHAPTER 4  CASE ANALYSIS\t79",
        111: "4.1. Case Overview and Analytical Focus\t79",
        112: "4.2. Formation of AI Utilization Training Requests and the Organizational Judgment Stage\t81",
        113: "4.3. Failure of Entry into Formal Decision-Making Procedures and the Emergence of Non-decision\t84",
        114: "4.4. Analysis of Judgment\u2013Decision-Making Conversion Discontinuity through Process Tracing\t86",
        115: "4.5. Analysis of Non-decision Types through Pattern Matching\t88",
        116: "4.6. Synthesis of Findings: Judgment without Decision\t92",
        117: "4.7. Interim Summary of Case Analysis Findings\t94",

        # === TOC Chapter 5 ===
        118: "CHAPTER 5  DISCUSSION: PROCESSUAL THEORETICAL ELABORATION OF THE SEIZING MECHANISM\t99",
        119: "5.1. Theoretical Re-positioning of the Empirical Findings\t99",
        120: "5.2. Structural Characteristics of the Failed Judgment\u2013Decision Transition\t101",
        121: "5.3. Processual Dynamics of the Non-decision State\t104",
        122: "5.4. Processual Breakdown in Seizing and the Absence of Institutional Triggers\t107",
        123: "5.5. Processual Extension and Reinterpretation of Existing Seizing Theory\t113",
        124: "5.6. Formalizing the Processual Mechanism of Judgment\u2013Decision Conversion Stagnation\t117",
        125: "5.7. Theoretical Implications: Dynamic Capabilities from a Processual Perspective\t122",
        126: "5.8. Summary\t124",

        # === TOC Chapter 6 ===
        127: "CHAPTER 6  CONCLUSION AND IMPLICATIONS\t127",
        128: "6.1. Summary of Research Questions and Research Approach\t127",
        129: "6.2. Structural Synthesis of Key Findings: Judgment\u2013Decision Stagnation Mechanism\t128",
        130: "6.3. Theoretical Implications: Processual Redefinition and Theoretical Extension of the Seizing Stage\t131",
        131: "6.4. Theoretical Generalization and Scope Conditions\t134",
        132: "6.5. Practical Implications\t136",
        133: "6.6. Limitations of the Study\t138",
        134: "6.7. Summary\t141",

        # === TOC Chapter 7 ===
        135: "CHAPTER 7  RECOMMENDATIONS AND DIRECTIONS FOR FUTURE RESEARCH\t144",
        136: "7.1. Theoretical Implications: Processual Re-conceptualization of Seizing Mechanisms\t144",
        137: "7.2. Practical Implications: Organizational Design Principles for Judgment\u2013Decision Conversion\t147",
        138: "7.3. Policy and Institutional Recommendations for Enabling Judgment\u2013Decision Transition\t150",
        139: "7.4. Methodological Recommendations: Studying Non-Events and Processual Discontinuities\t152",
        140: "7.5. Future Research Directions on Judgment\u2013Decision Transitions\t154",
        141: "7.6. Overall Contributions and Implications\t156",
        142: "Summary\t159",
        143: "Conclusion\t163",
        144: "Recommendations\t165",

        # === TOC back matter ===
        155: "REFERENCES CITED\t186",
        156: "Appendices\t190",

        # === Table list entries ===
        160: "<Table 3.1.>  List and Volume of Data\t60",
        162: "<Table 3.2.>  Correspondence between Appendices and Analysis Stages\t62",
        165: "<Table 3.3.>  Operational Indicators and Interpretive Memos for Judgment\u2013Decision Conversion Analysis\t64",

        # === Figure list entries ===
        184: "<Figure 1.1>  Overall Logical Flow and Analytical Framework of This Study\t29",
        186: "<Figure 4.1>  Structural Disconnection of Judgment\u2013Decision Conversion and the Seizing Non-initiation Mechanism\t92",
        188: "<Figure 4.2>  Processual Bottleneck Mechanism at the Seizing Stage Following Organizational AI Utilization Judgment\t97",
        190: "<Figure 5.1>  Structural Diagram of the Self-Reinforcing Mechanism of the Non-decision State\t110",
        192: "<Figure 5.2>  Processual Mechanism of Judgment\u2013Decision-Making Conversion Stagnation at the Seizing Stage\t118",
        194: "<Figure 5.3>  Judgment Life Cycle Framework and Conditional Conversion Pathways\t121",
        196: "<Figure 6.1>  The Judgment\u2013Decision Non-Initiation Process\t130",

        # === Dedication ===
        213: "This dissertation is dedicated to",
        214: "my beloved family,\nwho stood by me with quiet support and patience\nthrough every difficulty,\nenabling me to complete this work.",
        215: "Their unwavering patience and devotion\nwere the greatest strength\nthat sustained this academic journey to its end.",

        # === Acknowledgements ===
        219: "The journey to completing this dissertation was by no means a process that could be borne alone. I extend my deepest gratitude to the many mentors who shared in contemplating the academic and practical directions of this research and offered unstinting counsel until this study was completed as a doctoral dissertation.",
        220: "In particular, I express my sincere gratitude to Advisor Professor Jini Choi and Co-Advisor Professor Youngjun Choi, who guided me with consistent academic standards and insight throughout this research. Their meticulous guidance and probing questions played a decisive role in equipping this study with clearer problem consciousness and greater theoretical depth.",
        221: "I also extend my gratitude to OIKOS University for providing a stable research environment and scholarly community throughout the doctoral program. The systematic curriculum and academic atmosphere served as an important foundation for sustaining and deepening this research.",
        222: "I also wish to convey my gratitude to the fellow researchers who shared assistance and encouragement in diverse ways throughout the entire research process. Every moment and experience shared together served as a precious occasion for developing this research more reflectively.",
        223: "Henceforth, rather than the distinguished title of doctorate, I shall not forget the posture of a researcher who asks honestly and thinks deeply. I shall humbly walk the path of scholarship that benefits the world from the lowest places.",
    }

    # Apply all translations using XML-level clearing
    replaced = 0
    for i, para in enumerate(doc.paragraphs):
        if i >= 228:
            break
        if i in translations:
            clear_paragraph_content(para, translations[i])
            replaced += 1

    doc.save(working)
    print(f"Replaced {replaced} front matter paragraphs")

    # Verify
    doc2 = Document(working)
    kr_pattern = re.compile(r'[\uac00-\ud7af]')
    remaining = []
    for i, para in enumerate(doc2.paragraphs):
        if i >= 228:
            break
        if kr_pattern.search(para.text):
            remaining.append(f"P{i} [{para.style.name}]: {para.text[:80]}")

    print(f"Korean residuals: {len(remaining)}")
    if remaining:
        for r in remaining[:15]:
            print(f"  {r}")

    return replaced

if __name__ == '__main__':
    translate_front_matter()
