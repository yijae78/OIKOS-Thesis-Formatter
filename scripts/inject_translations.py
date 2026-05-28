"""Inject all translations from JSON into thesis_english.docx.
Replaces Korean body text (P228-P981) with English translations.
Preserves paragraph styles, formatting, and document structure."""
import json
import os
import sys
import shutil
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt

def inject_translations():
    base = os.path.join(os.path.dirname(__file__), '..')
    working = os.path.join(base, 'output', 'thesis_english.docx')
    trans_file = os.path.join(base, 'output', 'translations_all.json')

    if not os.path.exists(working):
        print(f"ERROR: Working copy not found: {working}")
        sys.exit(1)

    # Backup before injection
    backup_dir = os.path.join(base, 'backups')
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup = os.path.join(backup_dir, f'thesis_pre_inject_{ts}.docx')
    shutil.copy2(working, backup)
    print(f"Backup created: {os.path.basename(backup)}")

    # Load translations
    with open(trans_file, 'r', encoding='utf-8') as f:
        translations = json.load(f)
    print(f"Loaded {len(translations)} translations")

    # Open document
    doc = Document(working)
    total_paras = len(doc.paragraphs)
    print(f"Document has {total_paras} paragraphs")

    # Inject translations
    replaced = 0
    skipped_empty = 0
    errors = []

    for i, para in enumerate(doc.paragraphs):
        idx = str(i)
        if idx not in translations:
            continue

        new_text = translations[idx]

        try:
            if para.runs:
                # Strategy: put all text in first run, clear the rest
                # This preserves the first run's formatting (bold, italic, size)
                first_run = para.runs[0]

                # Store first run's key formatting
                was_bold = first_run.bold
                was_italic = first_run.italic
                font_size = first_run.font.size

                # Set translated text
                first_run.text = new_text

                # Set font to Times New Roman for the translated run
                first_run.font.name = 'Times New Roman'

                # Clear remaining runs
                for run in para.runs[1:]:
                    run.text = ""

                replaced += 1
            else:
                # No runs - add one with the translated text
                run = para.add_run(new_text)
                run.font.name = 'Times New Roman'
                replaced += 1

        except Exception as e:
            errors.append(f"P{i}: {str(e)}")

    print(f"\nInjection results:")
    print(f"  Replaced: {replaced}/{len(translations)}")
    if errors:
        print(f"  Errors: {len(errors)}")
        for err in errors[:5]:
            print(f"    {err}")

    # Save
    doc.save(working)
    print(f"\nSaved: {os.path.basename(working)}")

    # Quick verification
    print("\n--- Verification ---")
    doc2 = Document(working)
    verify_ok = 0
    verify_korean = 0
    korean_pattern = re.compile(r'[\uac00-\ud7af]')

    for i, para in enumerate(doc2.paragraphs):
        idx = str(i)
        if idx in translations:
            text = para.text.strip()
            if text:
                verify_ok += 1
                if korean_pattern.search(text):
                    verify_korean += 1

    print(f"  Paragraphs with content: {verify_ok}/{len(translations)}")
    print(f"  Korean residuals: {verify_korean}")

    # Sample check: show first 3 translated paragraphs
    print("\n--- Sample (first 3 body paragraphs) ---")
    count = 0
    for i, para in enumerate(doc2.paragraphs):
        idx = str(i)
        if idx in translations and i >= 228:
            text = para.text[:100]
            print(f"  P{i}: {text}...")
            count += 1
            if count >= 3:
                break

    return replaced

if __name__ == '__main__':
    replaced = inject_translations()
    print(f"\nDone. {replaced} paragraphs injected.")
