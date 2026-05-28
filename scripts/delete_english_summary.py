"""Delete ENGLISH SUMMARY sections (S29-S44) from thesis_english.docx.

The entire dissertation is now in English, so the ENGLISH SUMMARY that was
included as a required section in the Korean version is redundant.

After deletion:
  - S28 (Chapter 7) connects directly to what was S45 (REFERENCES)
  - Section count drops from 47 to 31
  - Orphaned header/footer relationship parts are cleaned up

Safety:
  - Creates a timestamped backup before any modification
  - Validates the result with python-docx after deletion
  - Reports paragraph and section counts for manual verification
"""

import sys
import shutil
from pathlib import Path
from datetime import datetime
from copy import deepcopy

# Ensure project root is on path
PROJECT_DIR = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_DIR))

from docx import Document
from docx.oxml.ns import qn
import lxml.etree as etree


# ── Configuration ────────────────────────────────────────────────────────
INPUT_FILE = PROJECT_DIR / "output" / "thesis_english.docx"
BACKUP_DIR = PROJECT_DIR / "backups"
BACKUP_DIR.mkdir(exist_ok=True)

# Sections to delete (0-based): S29 through S44 inclusive
DELETE_START = 29
DELETE_END = 44  # inclusive


def create_backup(src: Path) -> Path:
    """Create a timestamped backup of the source file."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_name = f"thesis_english_backup_{timestamp}_before_summary_delete.docx"
    backup_path = BACKUP_DIR / backup_name
    shutil.copy2(src, backup_path)
    size_mb = backup_path.stat().st_size / (1024 * 1024)
    print(f"[BACKUP] Created: {backup_path.name} ({size_mb:.1f} MB)")
    return backup_path


def get_section_boundaries(doc: Document) -> list[dict]:
    """Find XML element boundaries for each section.

    Returns list of dicts with:
        section_index, start_child, end_child, sectPr_location, sectPr_element
    """
    body = doc.element.body
    children = list(body)
    boundaries = []
    current_start = 0

    for idx, child in enumerate(children):
        if child.tag == qn('w:p'):
            pPr = child.find(qn('w:pPr'))
            if pPr is not None:
                sectPr = pPr.find(qn('w:sectPr'))
                if sectPr is not None:
                    boundaries.append({
                        'section_index': len(boundaries),
                        'start_child': current_start,
                        'end_child': idx,
                        'sectPr_location': 'paragraph',
                        'sectPr_element': sectPr,
                        'container': pPr,
                    })
                    current_start = idx + 1

    # Last section: body-level sectPr
    body_sectPr = body.find(qn('w:sectPr'))
    if body_sectPr is not None:
        boundaries.append({
            'section_index': len(boundaries),
            'start_child': current_start,
            'end_child': len(children) - 1,
            'sectPr_location': 'body',
            'sectPr_element': body_sectPr,
            'container': body,
        })

    return boundaries


def get_section_title(children, start, end):
    """Extract a readable title from section content."""
    for i in range(start, min(end + 1, len(children))):
        child = children[i]
        if child.tag == qn('w:p'):
            texts = []
            for r in child.findall(qn('w:r')):
                for t in r.findall(qn('w:t')):
                    if t.text:
                        texts.append(t.text)
            text = ''.join(texts).strip()
            if text:
                return text[:60]
    return "(empty)"


def collect_rids_for_sections(doc: Document, section_indices) -> set:
    """Collect all header/footer relationship IDs for given sections."""
    rids = set()
    for i in section_indices:
        sp = doc.sections[i]._sectPr
        for ref in sp.findall(qn('w:headerReference')) + sp.findall(qn('w:footerReference')):
            rid = ref.get(qn('r:id'))
            if rid:
                rids.add(rid)
    return rids


def delete_sections(doc: Document, start_sec: int, end_sec: int) -> dict:
    """Delete sections start_sec through end_sec (inclusive, 0-based).

    Strategy:
    1. Identify all body children (paragraphs/tables) in S29-S44
    2. Remove those children from the body
    3. Handle the sectPr boundary: S28's sectPr stays, S29-S44's sectPr elements
       are removed along with their host paragraphs
    4. Clean up orphaned header/footer relationship parts

    Returns a stats dict.
    """
    body = doc.element.body
    boundaries = get_section_boundaries(doc)
    children = list(body)

    total_sections_before = len(boundaries)
    total_children_before = len(children)

    # Validate range
    if start_sec < 0 or end_sec >= total_sections_before:
        raise ValueError(f"Section range {start_sec}-{end_sec} out of bounds (0-{total_sections_before - 1})")

    # ── Step 1: Identify children to remove ──────────────────────────
    first_child = boundaries[start_sec]['start_child']
    last_child = boundaries[end_sec]['end_child']

    print(f"\n[INFO] Sections to delete: S{start_sec:02d} - S{end_sec:02d} ({end_sec - start_sec + 1} sections)")
    print(f"[INFO] Body children to remove: indices {first_child} to {last_child} ({last_child - first_child + 1} elements)")

    # Show what we're deleting
    print(f"\n--- Content being deleted ---")
    for si in range(start_sec, end_sec + 1):
        b = boundaries[si]
        title = get_section_title(children, b['start_child'], b['end_child'])
        print(f"  S{si:02d} [{b['start_child']:4d}-{b['end_child']:4d}] {title}")
    print(f"--- End of deletion list ---\n")

    # ── Step 2: Collect rIds BEFORE deletion (sections won't exist after) ──
    rids_deleted = collect_rids_for_sections(doc, range(start_sec, end_sec + 1))
    # Also collect rIds for sections we are keeping
    keep_ranges = list(range(0, start_sec)) + list(range(end_sec + 1, total_sections_before))
    rids_kept = collect_rids_for_sections(doc, keep_ranges)
    orphaned_rids = rids_deleted - rids_kept
    print(f"[INFO] rIds in deleted sections: {len(rids_deleted)}, orphaned: {len(orphaned_rids)}")

    # ── Step 3: Collect elements to remove ───────────────────────────
    elements_to_remove = []
    for idx in range(first_child, last_child + 1):
        elements_to_remove.append(children[idx])

    print(f"[INFO] Removing {len(elements_to_remove)} body-level elements...")

    # ── Step 4: Remove elements from body ────────────────────────────
    removed_count = 0
    for elem in elements_to_remove:
        body.remove(elem)
        removed_count += 1

    print(f"[OK] Removed {removed_count} elements from document body")

    # ── Step 5: Clean up orphaned header/footer parts ────────────────
    cleaned_parts = 0
    if orphaned_rids:
        try:
            doc_part = doc.part
            rels = doc_part.rels
            for rid in orphaned_rids:
                if rid in rels:
                    del rels[rid]
                    cleaned_parts += 1
            print(f"[OK] Cleaned {cleaned_parts} orphaned relationship parts")
        except Exception as e:
            print(f"[WARN] Could not clean orphaned parts: {e}")
            print(f"       (Document will still work; Word will ignore unused parts)")

    # ── Step 6: Verify new structure ─────────────────────────────────
    new_boundaries = get_section_boundaries(doc)
    new_children = list(body)

    stats = {
        'sections_before': total_sections_before,
        'sections_after': len(new_boundaries),
        'sections_deleted': total_sections_before - len(new_boundaries),
        'children_before': total_children_before,
        'children_after': len(new_children),
        'children_removed': removed_count,
        'orphaned_rids_cleaned': cleaned_parts,
    }

    return stats


def verify_document(filepath: Path) -> bool:
    """Verify the document can be opened and has expected structure."""
    try:
        doc = Document(str(filepath))
        sec_count = len(doc.sections)
        para_count = len(doc.paragraphs)
        table_count = len(doc.tables)

        print(f"\n[VERIFY] Document opened successfully")
        print(f"  Sections:   {sec_count}")
        print(f"  Paragraphs: {para_count}")
        print(f"  Tables:     {table_count}")

        # Check that the section after Chapter 7 is REFERENCES
        boundaries = get_section_boundaries(doc)
        children = list(doc.element.body)

        # Find the section that was previously S45 (now should be around S29)
        found_refs = False
        for b in boundaries:
            title = get_section_title(children, b['start_child'], b['end_child'])
            if 'REFERENCES' in title.upper() or 'CITED' in title.upper():
                print(f"  REFERENCES found at S{b['section_index']:02d}: {title}")
                found_refs = True
                break

        if not found_refs:
            print(f"  [WARN] Could not find REFERENCES section")

        # Check no ENGLISH SUMMARY remains
        found_summary = False
        for b in boundaries:
            title = get_section_title(children, b['start_child'], b['end_child'])
            if 'ENGLISH SUMMARY' in title.upper():
                print(f"  [WARN] ENGLISH SUMMARY still present at S{b['section_index']:02d}")
                found_summary = True

        if not found_summary:
            print(f"  [OK] No ENGLISH SUMMARY sections remain")

        return True
    except Exception as e:
        print(f"\n[VERIFY FAIL] {e}")
        return False


def print_new_section_layout(filepath: Path):
    """Print the section layout after deletion for visual confirmation."""
    doc = Document(str(filepath))
    boundaries = get_section_boundaries(doc)
    children = list(doc.element.body)

    print(f"\n{'=' * 70}")
    print(f"NEW SECTION LAYOUT ({len(boundaries)} sections)")
    print(f"{'=' * 70}")
    for b in boundaries:
        title = get_section_title(children, b['start_child'], b['end_child'])
        loc = b['sectPr_location']
        elem_count = b['end_child'] - b['start_child'] + 1
        print(f"  S{b['section_index']:02d} [{b['start_child']:4d}-{b['end_child']:4d}] elems={elem_count:4d} {loc:10s} | {title}")
    print(f"{'=' * 70}")


def main():
    print("=" * 70)
    print("DELETE ENGLISH SUMMARY SECTIONS (S29-S44)")
    print("=" * 70)

    if not INPUT_FILE.exists():
        print(f"[ERROR] Input file not found: {INPUT_FILE}")
        sys.exit(1)

    # ── Step 0: Backup ───────────────────────────────────────────────
    backup_path = create_backup(INPUT_FILE)

    # ── Step 1: Load document ────────────────────────────────────────
    print(f"\n[INFO] Loading: {INPUT_FILE.name}")
    doc = Document(str(INPUT_FILE))
    print(f"[INFO] Loaded: {len(doc.sections)} sections, {len(doc.paragraphs)} paragraphs")

    # ── Step 2: Delete sections ──────────────────────────────────────
    stats = delete_sections(doc, DELETE_START, DELETE_END)

    print(f"\n--- Deletion Statistics ---")
    print(f"  Sections: {stats['sections_before']} -> {stats['sections_after']} (deleted {stats['sections_deleted']})")
    print(f"  Body children: {stats['children_before']} -> {stats['children_after']} (removed {stats['children_removed']})")
    print(f"  Orphaned parts cleaned: {stats['orphaned_rids_cleaned']}")

    expected_sections = stats['sections_before'] - (DELETE_END - DELETE_START + 1)
    if stats['sections_after'] != expected_sections:
        print(f"  [WARN] Expected {expected_sections} sections, got {stats['sections_after']}")
    else:
        print(f"  [OK] Section count matches expected: {expected_sections}")

    # ── Step 3: Save ─────────────────────────────────────────────────
    print(f"\n[INFO] Saving to: {INPUT_FILE.name}")
    doc.save(str(INPUT_FILE))
    size_mb = INPUT_FILE.stat().st_size / (1024 * 1024)
    print(f"[OK] Saved ({size_mb:.1f} MB)")

    # ── Step 4: Verify ───────────────────────────────────────────────
    if verify_document(INPUT_FILE):
        print(f"\n[SUCCESS] ENGLISH SUMMARY sections deleted successfully")
        print_new_section_layout(INPUT_FILE)
    else:
        print(f"\n[FAIL] Verification failed! Restoring from backup...")
        shutil.copy2(backup_path, INPUT_FILE)
        print(f"[RESTORED] Original file restored from: {backup_path.name}")
        sys.exit(1)

    print(f"\nBackup available at: {backup_path}")


if __name__ == '__main__':
    main()
