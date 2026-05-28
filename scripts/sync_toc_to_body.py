"""Sync TOC entries to match actual body heading text.
Also removes stale ENGLISH SUMMARY entries from TOC."""
import os
import re
import sys
import io
import shutil
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

BASE = os.path.join(os.path.dirname(__file__), '..')
WORKING = os.path.join(BASE, 'output', 'thesis_english.docx')


def clear_paragraph_content(para, new_text, font_name='Times New Roman'):
    """Clear all content and set new text (preserves pPr)."""
    elem = para._element
    for child in list(elem):
        if child.tag != qn('w:pPr'):
            elem.remove(child)
    if new_text:
        run = para.add_run(new_text)
        run.font.name = font_name


def main():
    # Backup
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup = os.path.join(BASE, 'backups', f'thesis_pre_tocsync_{ts}.docx')
    shutil.copy2(WORKING, backup)
    print(f"Backup: {os.path.basename(backup)}")

    doc = Document(WORKING)

    # === 1. Collect body headings by section number ===
    body_headings = {}
    heading_styles = ['*Heading: Main', '*Heading: 1st-Level', '*Heading: 2nd-Level',
                      'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4',
                      '*Heading: TOC']
    for i, p in enumerate(doc.paragraphs):
        if any(hs in p.style.name for hs in heading_styles):
            txt = p.text.strip().replace('\n', ' ').replace('  ', ' ')
            if txt and len(txt) < 300:
                # Extract section number
                m = re.match(r'^(\d+\.\d+\.?\d*\.?)\s', txt)
                if m:
                    key = m.group(1).rstrip('.')
                    body_headings[key] = txt

    # Also collect chapter headings
    for i, p in enumerate(doc.paragraphs):
        if '*Heading: Main' in p.style.name:
            txt = p.text.strip().replace('\n', ' ').replace('  ', ' ')
            m = re.match(r'^CHAPTER\s+(\d+)', txt)
            if m:
                # Format: "CHAPTER N  TITLE"
                ch_num = m.group(1)
                body_headings[f'ch{ch_num}'] = txt

    # Also find Summary/Conclusion/Recommendations headings (Chapter 7 sub-sections)
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if p.style.name in heading_styles or 'Heading' in p.style.name:
            if txt.startswith('Summary') and i > 500:
                body_headings['summary_end'] = txt
            elif txt.startswith('Conclusion') and i > 500:
                body_headings['conclusion_end'] = txt
            elif txt.startswith('Recommendations') and i > 500:
                body_headings['recommendations_end'] = txt

    print(f"Body headings collected: {len(body_headings)}")

    # === 2. Find TOC range ===
    toc_start = toc_end = None
    for i, p in enumerate(doc.paragraphs):
        if 'TABLE OF CONTENTS' in p.text and ('Heading' in p.style.name or 'Main' in p.style.name):
            toc_start = i + 1
        if toc_start and i > toc_start and 'toc' not in p.style.name.lower():
            if p.text.strip() and i > toc_start + 5:
                toc_end = i
                break

    print(f"TOC range: P{toc_start} to P{toc_end}")

    # === 3. Process TOC entries ===
    updated = 0
    removed = 0
    english_summary_zone = False

    for i in range(toc_start, toc_end):
        p = doc.paragraphs[i]
        if 'toc' not in p.style.name.lower():
            continue
        txt = p.text.strip()
        if not txt:
            continue

        # Parse TOC entry: "title\tpagenum"
        parts = txt.rsplit('\t', 1)
        toc_title = parts[0].strip()
        toc_pgnum = parts[1].strip() if len(parts) > 1 else ''

        # Check for ENGLISH SUMMARY zone
        if 'ENGLISH SUMMARY' in toc_title:
            english_summary_zone = True

        if english_summary_zone:
            # Remove this entry (ENGLISH SUMMARY section was deleted)
            clear_paragraph_content(p, '')
            removed += 1
            continue

        # Try to match by section number
        m = re.match(r'^(\d+\.\d+\.?\d*\.?)\s', toc_title)
        if m:
            key = m.group(1).rstrip('.')
            if key in body_headings:
                body_title = body_headings[key]
                new_entry = f"{body_title}\t{toc_pgnum}"
                if body_title != toc_title:
                    clear_paragraph_content(p, new_entry)
                    updated += 1
                    print(f"  SYNC P{i}: {toc_title[:50]}")
                    print(f"       --> {body_title[:50]}")

        # Match chapter headings
        m_ch = re.match(r'^CHAPTER\s+(\d+)', toc_title)
        if m_ch:
            ch_num = m_ch.group(1)
            key = f'ch{ch_num}'
            if key in body_headings:
                body_ch = body_headings[key]
                # TOC format: "CHAPTER N  TITLE\tpgnum"
                # Body format: "CHAPTER N  TITLE" (newlines replaced with spaces)
                if body_ch.replace('  ', ' ') != toc_title.replace('  ', ' '):
                    new_entry = f"{body_ch}\t{toc_pgnum}"
                    clear_paragraph_content(p, new_entry)
                    updated += 1
                    print(f"  SYNC P{i}: {toc_title[:50]}")
                    print(f"       --> {body_ch[:50]}")

        # Match Summary/Conclusion/Recommendations (end of Chapter 7)
        if toc_title.startswith('Summary') and toc_pgnum:
            if 'summary_end' in body_headings:
                new_entry = f"{body_headings['summary_end']}\t{toc_pgnum}"
                clear_paragraph_content(p, new_entry)
                updated += 1
        elif toc_title.startswith('Conclusion') and toc_pgnum:
            if 'conclusion_end' in body_headings:
                new_entry = f"{body_headings['conclusion_end']}\t{toc_pgnum}"
                clear_paragraph_content(p, new_entry)
                updated += 1
        elif toc_title.startswith('Recommendations') and toc_pgnum:
            if 'recommendations_end' in body_headings:
                new_entry = f"{body_headings['recommendations_end']}\t{toc_pgnum}"
                clear_paragraph_content(p, new_entry)
                updated += 1

    print(f"\nResults: {updated} entries synced, {removed} ENGLISH SUMMARY entries removed")

    # Save
    doc.save(WORKING)
    print(f"Saved: {os.path.basename(WORKING)}")

    # === 4. Verify ===
    print("\n--- Verification ---")
    doc2 = Document(WORKING)
    mismatch_count = 0

    # Re-collect body headings
    bh2 = {}
    for i, p in enumerate(doc2.paragraphs):
        if any(hs in p.style.name for hs in heading_styles):
            txt = p.text.strip().replace('\n', ' ')
            m = re.match(r'^(\d+\.\d+\.?\d*\.?)\s', txt)
            if m:
                bh2[m.group(1).rstrip('.')] = txt

    for i in range(toc_start, min(toc_end, len(doc2.paragraphs))):
        p = doc2.paragraphs[i]
        if 'toc' not in p.style.name.lower():
            continue
        txt = p.text.strip()
        if not txt:
            continue
        parts = txt.rsplit('\t', 1)
        toc_title = parts[0].strip()

        m = re.match(r'^(\d+\.\d+\.?\d*\.?)\s', toc_title)
        if m:
            key = m.group(1).rstrip('.')
            if key in bh2 and bh2[key] != toc_title:
                print(f"  STILL MISMATCH P{i}: TOC={toc_title[:50]} BODY={bh2[key][:50]}")
                mismatch_count += 1

    # Check no ENGLISH SUMMARY entries remain
    es_remaining = 0
    for i in range(toc_start, toc_end):
        p = doc2.paragraphs[i]
        if 'ENGLISH SUMMARY' in p.text:
            es_remaining += 1

    print(f"  Remaining mismatches: {mismatch_count}")
    print(f"  ENGLISH SUMMARY entries: {es_remaining}")
    print(f"  Total TOC entries: {sum(1 for i in range(toc_start, toc_end) if doc2.paragraphs[i].text.strip())}")


if __name__ == '__main__':
    main()
