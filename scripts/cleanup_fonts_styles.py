"""Phase 9: Font and style cleanup for English thesis.
- All Korean fonts (맑은 고딕, 바탕체) → Times New Roman
- Theme fonts → Times New Roman
- eastAsia language → en-US
- Paragraph run fonts → Times New Roman"""
import os
import shutil
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from lxml import etree

def cleanup_fonts():
    base = os.path.join(os.path.dirname(__file__), '..')
    working = os.path.join(base, 'output', 'thesis_english.docx')

    # Backup
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup = os.path.join(base, 'backups', f'thesis_pre_fonts_{ts}.docx')
    shutil.copy2(working, backup)
    print(f"Backup: {os.path.basename(backup)}")

    doc = Document(working)

    # === 1. Fix theme fonts (Major/Minor) ===
    try:
        theme_part = doc.part.part_related_by('http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme')
    except KeyError:
        theme_part = None
    if theme_part:
        theme_xml = theme_part.blob
        tree = etree.fromstring(theme_xml)
        ns = {
            'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'
        }
        # Find majorFont and minorFont, set latin to Times New Roman
        for font_scheme in tree.findall('.//a:majorFont', ns):
            latin = font_scheme.find('a:latin', ns)
            if latin is not None:
                latin.set('typeface', 'Times New Roman')
            ea = font_scheme.find('a:ea', ns)
            if ea is not None:
                ea.set('typeface', 'Times New Roman')
        for font_scheme in tree.findall('.//a:minorFont', ns):
            latin = font_scheme.find('a:latin', ns)
            if latin is not None:
                latin.set('typeface', 'Times New Roman')
            ea = font_scheme.find('a:ea', ns)
            if ea is not None:
                ea.set('typeface', 'Times New Roman')
        theme_part._blob = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
        print("Theme fonts updated to Times New Roman")

    # === 2. Fix document default fonts ===
    styles_element = doc.styles.element
    doc_defaults = styles_element.find(qn('w:docDefaults'))
    if doc_defaults is not None:
        rpr_default = doc_defaults.find(f'.//{qn("w:rPr")}')
        if rpr_default is not None:
            rfonts = rpr_default.find(qn('w:rFonts'))
            if rfonts is not None:
                for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                    rfonts.set(qn(attr), 'Times New Roman')
                # Remove theme font references
                for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
                    if qn(attr) in rfonts.attrib:
                        del rfonts.attrib[qn(attr)]
                print("Document default fonts set to Times New Roman")

    # === 3. Fix all style definitions ===
    styles_changed = 0
    for style in doc.styles:
        try:
            style_elem = style.element
            rpr = style_elem.find(qn('w:rPr'))
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    changed = False
                    for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                        val = rfonts.get(qn(attr), '')
                        if val and val != 'Times New Roman' and 'Symbol' not in val:
                            rfonts.set(qn(attr), 'Times New Roman')
                            changed = True
                    # Remove theme references
                    for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
                        if qn(attr) in rfonts.attrib:
                            del rfonts.attrib[qn(attr)]
                            changed = True
                    if changed:
                        styles_changed += 1
        except Exception:
            pass
    print(f"Style definitions updated: {styles_changed}")

    # === 4. Fix all paragraph runs ===
    runs_changed = 0
    for para in doc.paragraphs:
        for run in para.runs:
            rpr = run._element.find(qn('w:rPr'))
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                        val = rfonts.get(qn(attr), '')
                        if val and val != 'Times New Roman' and 'Symbol' not in val:
                            rfonts.set(qn(attr), 'Times New Roman')
                            runs_changed += 1
                    # Remove theme references from runs
                    for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
                        if qn(attr) in rfonts.attrib:
                            del rfonts.attrib[qn(attr)]

    # Also fix table cell runs
    table_runs = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        rpr = run._element.find(qn('w:rPr'))
                        if rpr is not None:
                            rfonts = rpr.find(qn('w:rFonts'))
                            if rfonts is not None:
                                for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                                    val = rfonts.get(qn(attr), '')
                                    if val and val != 'Times New Roman' and 'Symbol' not in val:
                                        rfonts.set(qn(attr), 'Times New Roman')
                                        table_runs += 1
                                for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
                                    if qn(attr) in rfonts.attrib:
                                        del rfonts.attrib[qn(attr)]

    print(f"Paragraph runs font-fixed: {runs_changed}")
    print(f"Table runs font-fixed: {table_runs}")

    # === 5. Fix header/footer runs ===
    hf_runs = 0
    for section in doc.sections:
        for hf in [section.header, section.footer,
                    section.first_page_header, section.first_page_footer]:
            try:
                for para in hf.paragraphs:
                    for run in para.runs:
                        rpr = run._element.find(qn('w:rPr'))
                        if rpr is not None:
                            rfonts = rpr.find(qn('w:rFonts'))
                            if rfonts is not None:
                                for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                                    val = rfonts.get(qn(attr), '')
                                    if val and val != 'Times New Roman' and 'Symbol' not in val:
                                        rfonts.set(qn(attr), 'Times New Roman')
                                        hf_runs += 1
                                for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
                                    if qn(attr) in rfonts.attrib:
                                        del rfonts.attrib[qn(attr)]
            except Exception:
                pass
    print(f"Header/footer runs font-fixed: {hf_runs}")

    # === 6. Fix language settings ===
    body_xml = doc.element
    # Find all w:lang elements and set to en-US
    lang_count = 0
    for lang_elem in body_xml.iter(qn('w:lang')):
        for attr in ['w:val', 'w:eastAsia', 'w:bidi']:
            if qn(attr) in lang_elem.attrib:
                old = lang_elem.get(qn(attr))
                if old in ['ko-KR', 'zh-CN', 'ja-JP']:
                    lang_elem.set(qn(attr), 'en-US')
                    lang_count += 1
    print(f"Language settings fixed: {lang_count}")

    # Save
    doc.save(working)
    print(f"\nSaved: {os.path.basename(working)}")

    # Verify
    doc2 = Document(working)
    print(f"\nVerification: Document opens successfully")
    print(f"  Sections: {len(doc2.sections)}")
    print(f"  Paragraphs: {len(doc2.paragraphs)}")

    # Check for remaining non-TNR fonts
    non_tnr = set()
    for para in doc2.paragraphs[:50]:
        for run in para.runs:
            if run.font.name and run.font.name != 'Times New Roman':
                non_tnr.add(run.font.name)
    if non_tnr:
        print(f"  Non-TNR fonts in first 50 paras: {non_tnr}")
    else:
        print(f"  All fonts in first 50 paras: Times New Roman")

if __name__ == '__main__':
    cleanup_fonts()
